import io
import uuid
from decimal import Decimal
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import text

from backend.database import get_db

router = APIRouter(prefix="/api/v1/sales/invoices", tags=["Sales Exports"])

@router.get("/export/excel")
def export_invoices_excel(db: Session = Depends(get_db)):
    """Exports all customer invoices into an Excel spreadsheet with styling."""
    query = text("""
        SELECT invoice_number, customer_name, customer_tax_number, issue_date, status, subtotal, vat_amount, grand_total 
        FROM customer_invoices 
        ORDER BY created_at DESC
    """)
    rows = db.execute(query).fetchall()
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Customer Invoices"

        headers = ["Invoice Number", "Customer Name", "Customer VAT", "Issue Date", "Status", "Subtotal (SAR)", "VAT (15%)", "Grand Total (SAR)"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r in rows:
            m = r._mapping
            issue_date = m["issue_date"]
            ws.append([
                m["invoice_number"] or "",
                m["customer_name"] or "",
                m["customer_tax_number"] or "N/A",
                str(issue_date.strftime("%Y-%m-%d %H:%M")) if hasattr(issue_date, 'strftime') and issue_date else (str(issue_date) if issue_date else ""),
                m["status"] or "",
                float(m["subtotal"]) if m["subtotal"] is not None else 0.0,
                float(m["vat_amount"]) if m["vat_amount"] is not None else 0.0,
                float(m["grand_total"]) if m["grand_total"] is not None else 0.0,
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_idx = col[0].column
            if col_idx is None:
                continue
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
    except Exception as exc:
        import pandas as pd
        data = [dict(r._mapping) for r in rows]
        df = pd.DataFrame(data)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Customer Invoices')
        output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=customer_invoices_report.xlsx"}
    )

@router.get("/{invoice_id}/pdf")
def download_invoice_pdf(invoice_id: str, db: Session = Depends(get_db)):
    """Generates a professional ZATCA-compliant PDF invoice with embedded QR code."""
    invoice = None
    try:
        inv_uuid = uuid.UUID(invoice_id)
        invoice = db.execute(
            text("SELECT * FROM customer_invoices WHERE id = :id"), {"id": inv_uuid}
        ).fetchone()
    except ValueError:
        invoice = db.execute(
            text("SELECT * FROM customer_invoices WHERE invoice_number = :num"), {"num": invoice_id}
        ).fetchone()

    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")

    inv = invoice._mapping

    zatca_log = db.execute(
        text("SELECT qr_code_payload, invoice_hash FROM zatca_logs WHERE invoice_id = :id ORDER BY created_at DESC LIMIT 1"),
        {"id": inv["id"]}
    ).fetchone()

    # Generate PDF using ReportLab
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    import qrcode

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    story = []
    styles = getSampleStyleSheet()

    # Header
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=20, textColor=colors.HexColor('#1e293b'), spaceAfter=6)
    story.append(Paragraph("Tax Invoice / فاتورة ضريبية", title_style))
    story.append(Paragraph(f"<b>Invoice Number:</b> {inv['invoice_number']}", styles['Normal']))
    story.append(Paragraph(f"<b>Issue Date:</b> {inv['issue_date']}", styles['Normal']))
    story.append(Spacer(1, 12))

    # Party Info Table
    data = [
        [Paragraph("<b>Seller:</b> Myon Economic Co Ltd<br/><b>VAT ID:</b> 310123456700003", styles['Normal']),
         Paragraph(f"<b>Customer:</b> {inv['customer_name']}<br/><b>Customer VAT:</b> {inv['customer_tax_number'] or 'N/A'}", styles['Normal'])]
    ]
    t = Table(data, colWidths=[270, 270])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('PADDING', (0,0), (-1,-1), 10),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
    ]))
    story.append(t)
    story.append(Spacer(1, 16))

    # Line Items / Totals Table
    subtotal_val = float(inv["subtotal"]) if inv["subtotal"] is not None else 0.0
    vat_val = float(inv["vat_amount"]) if inv["vat_amount"] is not None else 0.0
    grand_val = float(inv["grand_total"]) if inv["grand_total"] is not None else 0.0

    totals_data = [
        ["Description", "Amount (SAR)"],
        [f"{inv['customer_name']} - Services rendered", f"{subtotal_val:.2f}"],
        ["Subtotal (Net)", f"{subtotal_val:.2f}"],
        ["VAT (15%)", f"{vat_val:.2f}"],
        ["Grand Total (SAR)", f"{grand_val:.2f}"]
    ]
    t2 = Table(totals_data, colWidths=[380, 160])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t2)

    # QR Code & Cryptographic Stamp
    zl = zatca_log._mapping if zatca_log else None
    if zl and zl["qr_code_payload"]:
        try:
            import qrcode.image.pil
            qr_img = qrcode.make(zl["qr_code_payload"], image_factory=qrcode.image.pil.PilImage)
            qr_buffer = io.BytesIO()
            qr_img.save(qr_buffer, format="PNG")
            qr_buffer.seek(0)
            story.append(Spacer(1, 14))
            story.append(RLImage(qr_buffer, width=110, height=110))
            if zl["invoice_hash"]:
                story.append(Spacer(1, 4))
                story.append(Paragraph(f"<b>ZATCA Hash:</b> <font size=7 color='#64748b'>{zl['invoice_hash']}</font>", styles['Normal']))
        except Exception:
            pass
    
    doc.build(story)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={inv['invoice_number']}.pdf"}
    )